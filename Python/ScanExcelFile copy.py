import os
from openpyxl import load_workbook

input_directory = r'\\fsval\archivio$\Docsvps\SCHEDE TECNICHE FORNITORI\LABORATORIO DAMIANI'  # <-- Modifica questo percorso
output_directory = r'c:\Users\rizzesim\OneDrive - CASA DAMIANI SPA\Documenti\Workspace Visual\Python'  # <-- Modifica questo percorso

def get_unique_filepath(path):
    base, ext = os.path.splitext(path)
    counter = 1
    candidate = path
    while os.path.exists(candidate):
        candidate = f"{base}_{counter}{ext}"
        counter += 1
    return candidate

os.makedirs(output_directory, exist_ok=True)
output_file = get_unique_filepath(os.path.join(output_directory, 'riepilogo_fogli.txt'))

with open(output_file, 'w', encoding='utf-8') as out:
    for root, dirs, files in os.walk(input_directory):
        for fname in files:
            # ignora file temporanei di Excel
            if fname.startswith('~$'):
                continue
            if fname.lower().endswith(('.xlsx', '.xlsm', '.xls')):
                path = os.path.join(root, fname)
                try:
                    wb = load_workbook(filename=path, read_only=True, data_only=True)
                    sheets = wb.sheetnames
                    out.write(f"File: {path}\n")
                    for s in sheets:
                        out.write(f"  - {s}\n")
                    out.write("\n")
                except Exception as e:
                    out.write(f"File: {path} - ERRORE: {e}\n\n")